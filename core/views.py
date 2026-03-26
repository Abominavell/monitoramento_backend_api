from datetime import date, datetime

from django.contrib.auth import authenticate
from django.db.models import Q
from rest_framework import permissions, status, viewsets
from rest_framework.decorators import api_view, permission_classes
from rest_framework.decorators import action
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.response import Response
from openpyxl import load_workbook

from .models import AcaoSocial, Colaborador, DataProjeto, Inscricao, Projeto, Setor
from .serializers import (
    AcaoSocialSerializer,
    ColaboradorSerializer,
    DataProjetoSerializer,
    InscricaoSerializer,
    ProjetoSerializer,
    SetorSerializer,
)
from .services import sync_colaborador_sede_user


class PublicReadAdminWrite(permissions.BasePermission):
    def has_permission(self, request, view):
        if request.method in permissions.SAFE_METHODS:
            return True
        return bool(request.user and request.user.is_authenticated)


class SetorViewSet(viewsets.ModelViewSet):
    queryset = Setor.objects.all()
    serializer_class = SetorSerializer
    permission_classes = [PublicReadAdminWrite]


class ColaboradorViewSet(viewsets.ModelViewSet):
    queryset = Colaborador.objects.select_related("setor").all()
    serializer_class = ColaboradorSerializer
    permission_classes = [PublicReadAdminWrite]
    parser_classes = [JSONParser, MultiPartParser, FormParser]

    def get_permissions(self):
        if self.request.method == "POST" and getattr(self, "action", None) == "create":
            return [permissions.AllowAny()]
        return [permission() for permission in self.permission_classes]

    def get_queryset(self):
        qs = super().get_queryset()
        p = self.request.query_params

        ativo = p.get("ativo")
        if ativo is None:
            qs = qs.filter(ativo=True)
        else:
            qs = qs.filter(ativo=ativo.lower() == "true")

        setor_id = p.get("setor_id")
        if setor_id:
            qs = qs.filter(setor_id=setor_id)

        nome = p.get("nome")
        if nome:
            qs = qs.filter(nome__iexact=nome)

        is_externo = p.get("is_externo")
        if is_externo is not None:
            qs = qs.filter(is_externo=is_externo.lower() == "true")

        ordering = p.get("ordering")
        if ordering:
            qs = qs.order_by(ordering)
        return qs

    @staticmethod
    def _parse_birth_date(raw_value):
        if isinstance(raw_value, datetime):
            return raw_value.date()
        if isinstance(raw_value, date):
            return raw_value
        if isinstance(raw_value, str):
            value = raw_value.strip()
            if not value:
                return None
            for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
                try:
                    return datetime.strptime(value, fmt).date()
                except ValueError:
                    continue
        return None

    @action(
        detail=False,
        methods=["post"],
        permission_classes=[permissions.IsAuthenticated],
        url_path="importar_excel",
    )
    def importar_excel(self, request):
        excel_file = request.FILES.get("file")
        if not excel_file:
            return Response({"detail": "Arquivo não enviado."}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb = load_workbook(excel_file, data_only=True)
            ws = wb.active
        except Exception:
            return Response({"detail": "Não foi possível ler o arquivo Excel."}, status=status.HTTP_400_BAD_REQUEST)

        header_map = {}
        for idx, cell in enumerate(ws[1], start=1):
            val = str(cell.value or "").strip().upper()
            if val:
                header_map[val] = idx

        required_headers = ["CPF", "NOME", "DATA NASCIMENTO", "SETOR"]
        missing = [h for h in required_headers if h not in header_map]
        if missing:
            return Response(
                {"detail": f"Colunas obrigatórias ausentes: {', '.join(missing)}"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        created = 0
        updated = 0
        skipped = 0
        errors = []
        cpfs_importados = set()

        for row_num in range(2, ws.max_row + 1):
            raw_cpf = ws.cell(row=row_num, column=header_map["CPF"]).value
            raw_nome = ws.cell(row=row_num, column=header_map["NOME"]).value
            raw_data_nasc = ws.cell(row=row_num, column=header_map["DATA NASCIMENTO"]).value
            raw_setor = ws.cell(row=row_num, column=header_map["SETOR"]).value

            cpf = "".join(ch for ch in str(raw_cpf or "") if ch.isdigit())
            nome = str(raw_nome or "").strip()
            setor_nome = str(raw_setor or "").strip()
            data_nascimento = self._parse_birth_date(raw_data_nasc)

            if not cpf and not nome and not setor_nome and not raw_data_nasc:
                continue

            if len(cpf) != 11 or not nome or not setor_nome or not data_nascimento:
                skipped += 1
                errors.append(f"Linha {row_num}: dados inválidos/incompletos.")
                continue

            cpfs_importados.add(cpf)
            setor = Setor.objects.filter(nome__iexact=setor_nome).first()
            if not setor:
                setor = Setor.objects.create(nome=setor_nome)

            try:
                colab = Colaborador.objects.filter(cpf=cpf).first()
                if colab:
                    colab.nome = nome
                    colab.setor = setor
                    colab.data_nascimento = data_nascimento
                    colab.is_externo = False
                    colab.ativo = True
                    colab.save()
                    updated += 1
                else:
                    Colaborador.objects.create(
                        cpf=cpf,
                        nome=nome,
                        setor=setor,
                        data_nascimento=data_nascimento,
                        is_externo=False,
                        ativo=True,
                    )
                    created += 1
            except Exception as exc:
                skipped += 1
                errors.append(f"Linha {row_num}: {str(exc)}")

        desativados = 0
        if cpfs_importados:
            desativados = Colaborador.objects.filter(is_externo=False, ativo=True).exclude(cpf__in=cpfs_importados).update(ativo=False)

        return Response(
            {
                "ok": True,
                "created": created,
                "updated": updated,
                "skipped": skipped,
                "deactivated": desativados,
                "message": "Importação concluída.",
                "errors": errors[:20],
            }
        )


class ProjetoViewSet(viewsets.ModelViewSet):
    queryset = Projeto.objects.all()
    serializer_class = ProjetoSerializer
    permission_classes = [PublicReadAdminWrite]

    def get_queryset(self):
        qs = super().get_queryset()
        p = self.request.query_params

        ativo = p.get("ativo")
        if ativo is not None:
            qs = qs.filter(ativo=ativo.lower() == "true")

        search = p.get("search")
        if search:
            qs = qs.filter(Q(titulo__icontains=search) | Q(descricao__icontains=search))

        ordering = p.get("ordering")
        if ordering:
            qs = qs.order_by(ordering)

        return qs


class DataProjetoViewSet(viewsets.ModelViewSet):
    queryset = DataProjeto.objects.select_related("projeto").all()
    serializer_class = DataProjetoSerializer
    permission_classes = [PublicReadAdminWrite]

    def get_queryset(self):
        qs = super().get_queryset()
        p = self.request.query_params

        projeto_id = p.get("projeto_id")
        if projeto_id:
            qs = qs.filter(projeto_id=projeto_id)

        ativo = p.get("ativo")
        if ativo is not None:
            qs = qs.filter(ativo=ativo.lower() == "true")

        data_gte = p.get("data_evento__gte")
        if data_gte:
            qs = qs.filter(data_evento__gte=data_gte)

        data_lte = p.get("data_evento__lte")
        if data_lte:
            qs = qs.filter(data_evento__lte=data_lte)

        ordering = p.get("ordering")
        if ordering:
            qs = qs.order_by(ordering)

        return qs


class AcaoSocialViewSet(viewsets.ModelViewSet):
    queryset = AcaoSocial.objects.all()
    serializer_class = AcaoSocialSerializer
    permission_classes = [PublicReadAdminWrite]

    def get_queryset(self):
        qs = super().get_queryset()
        p = self.request.query_params

        ativo = p.get("ativo")
        if ativo is not None:
            qs = qs.filter(ativo=ativo.lower() == "true")

        data_gte = p.get("data_evento__gte")
        if data_gte:
            qs = qs.filter(data_evento__gte=data_gte)

        data_lte = p.get("data_evento__lte")
        if data_lte:
            qs = qs.filter(data_evento__lte=data_lte)

        ordering = p.get("ordering")
        if ordering:
            qs = qs.order_by(ordering)
        return qs


class InscricaoViewSet(viewsets.ModelViewSet):
    queryset = Inscricao.objects.select_related("projeto", "data_projeto", "colaborador", "colaborador__setor").all()
    serializer_class = InscricaoSerializer
    permission_classes = [PublicReadAdminWrite]

    def get_permissions(self):
        if self.request.method == "POST":
            return [permissions.AllowAny()]
        return [permission() for permission in self.permission_classes]

    @action(
        detail=False,
        methods=["post"],
        permission_classes=[permissions.IsAuthenticated],
        url_path="reset_participacoes",
    )
    def reset_participacoes(self, request):
        """
        Reseta todas as participações confirmadas, voltando confirmado_presenca para False.
        """
        updated_count = Inscricao.objects.filter(confirmado_presenca=True).update(confirmado_presenca=False)
        return Response(
            {
                "ok": True,
                "updated": updated_count,
                "message": "Participações confirmadas resetadas com sucesso.",
            }
        )

    def get_queryset(self):
        qs = super().get_queryset()
        p = self.request.query_params

        projeto_id = p.get("projeto_id")
        if projeto_id:
            qs = qs.filter(projeto_id=projeto_id)

        data_projeto_id = p.get("data_projeto_id")
        if data_projeto_id:
            qs = qs.filter(data_projeto_id=data_projeto_id)

        data_evento_gte = p.get("data_evento__gte")
        if data_evento_gte:
            qs = qs.filter(data_projeto__data_evento__gte=data_evento_gte)

        data_evento_lte = p.get("data_evento__lte")
        if data_evento_lte:
            qs = qs.filter(data_projeto__data_evento__lte=data_evento_lte)

        colaborador_id = p.get("colaborador_id")
        if colaborador_id:
            qs = qs.filter(colaborador_id=colaborador_id)

        confirmado = p.get("confirmado_presenca")
        if confirmado is not None:
            qs = qs.filter(confirmado_presenca=confirmado.lower() == "true")

        created_gte = p.get("created_at__gte")
        if created_gte:
            qs = qs.filter(created_at__gte=created_gte)

        created_lt = p.get("created_at__lt")
        if created_lt:
            qs = qs.filter(created_at__lt=created_lt)

        colaborador_externo = p.get("colaborador__is_externo")
        if colaborador_externo is not None:
            qs = qs.filter(colaborador__is_externo=colaborador_externo.lower() == "true")

        search = p.get("search")
        if search:
            qs = qs.filter(
                Q(colaborador__nome__icontains=search)
                | Q(projeto__titulo__icontains=search)
            )

        ordering = p.get("ordering")
        if ordering:
            qs = qs.order_by(ordering)
        return qs


@api_view(["POST"])
@permission_classes([permissions.AllowAny])
def colaborador_login(request):
    cpf = (request.data.get("cpf") or "").strip()
    senha = (request.data.get("senha") or "").strip()
    cpf_digits = "".join(ch for ch in cpf if ch.isdigit())

    if not cpf or not senha:
        return Response(
            {"detail": "CPF e senha são obrigatórios."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    cpf_options = {cpf}
    if cpf_digits:
        cpf_options.add(cpf_digits)
    colaborador = Colaborador.objects.select_related("setor").filter(cpf__in=cpf_options, is_externo=False, ativo=True).first()
    if not colaborador:
        return Response({"detail": "Credenciais inválidas."}, status=status.HTTP_401_UNAUTHORIZED)

    sync_colaborador_sede_user(colaborador)
    user = authenticate(request, username=colaborador.cpf, password=senha)
    if not user:
        return Response({"detail": "Credenciais inválidas."}, status=status.HTTP_401_UNAUTHORIZED)

    return Response(
        {
            "ok": True,
            "colaborador_id": str(colaborador.id),
            "nome": colaborador.nome,
            "setor_id": str(colaborador.setor_id),
        }
    )
